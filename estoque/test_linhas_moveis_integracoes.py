from io import BytesIO

from django.contrib.auth.models import User
from django.core.exceptions import PermissionDenied
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from compras.models import CapacidadeCatalogoProdutoEmpresa, CatalogoProdutoEmpresa
from estoque.models import (
    Base,
    CategoriaEquipamentoEmpresa,
    Comunicado,
    CredencialLinhaMovel,
    Empresa,
    Equipamento,
    Historico,
    LinhaMovel,
    Modulo,
    ModuloEmpresa,
    OperadoraMovel,
    Perfil,
    Produto,
    VinculoLinhaEquipamento,
)
from estoque.services.assistente_operacional_service import AssistenteOperacionalService
from estoque.services.linhas_moveis_service import LinhasMoveisService


class LinhasMoveisIntegracoesSegurasTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.empresa = Empresa.objects.create(nome='Empresa integrações móveis')
        cls.empresa_externa = Empresa.objects.create(nome='Empresa externa integrações')
        cls.base = Base.objects.create(nome='Base integrações móveis', empresa=cls.empresa)
        cls.base_externa = Base.objects.create(
            nome='Base externa integrações',
            empresa=cls.empresa_externa,
        )

        cls.admin = User.objects.create_user('admin_integracoes_moveis')
        cls.admin_externo = User.objects.create_user('admin_externo_integracoes')
        cls.gestor_base = User.objects.create_user('gestor_base_integracoes')
        Perfil.objects.update_or_create(
            user=cls.admin,
            defaults={'empresa': cls.empresa, 'role': Perfil.Role.ADMIN},
        )
        Perfil.objects.update_or_create(
            user=cls.admin_externo,
            defaults={'empresa': cls.empresa_externa, 'role': Perfil.Role.ADMIN},
        )
        perfil_gestor, _ = Perfil.objects.update_or_create(
            user=cls.gestor_base,
            defaults={'empresa': cls.empresa, 'role': Perfil.Role.GESTOR},
        )
        perfil_gestor.regionais.add(cls.base)

        for codigo, nome in (
            (Modulo.Codigo.EQUIPAMENTOS, 'Equipamentos'),
            (Modulo.Codigo.TORY, 'Tory'),
        ):
            modulo, _ = Modulo.objects.update_or_create(
                codigo=codigo,
                defaults={'nome': nome, 'ativo': True},
            )
            ModuloEmpresa.objects.update_or_create(
                empresa=cls.empresa,
                modulo=modulo,
                defaults={'habilitado': True},
            )

        cls.produto = Produto.objects.create(
            codigo='PROD-INTEGRACAO-MOVEL',
            descricao='Coletor integração móvel',
            fabricante='Fabricante',
            modelo='Modelo móvel',
            categoria='Coletores integração',
        )
        CategoriaEquipamentoEmpresa.objects.create(
            empresa=cls.empresa,
            nome=cls.produto.categoria,
            aliases=['coletor', 'coletores'],
        )
        catalogo = CatalogoProdutoEmpresa.objects.create(
            empresa=cls.empresa,
            produto=cls.produto,
        )
        CapacidadeCatalogoProdutoEmpresa.objects.create(
            catalogo=catalogo,
            codigo=CapacidadeCatalogoProdutoEmpresa.CONECTIVIDADE_MOVEL,
        )

        cls.equipamento = Equipamento.objects.create(
            produto=cls.produto,
            regional=cls.base,
            numero_serie='SER-INTEGRACAO-MOVEL',
            patrimonio='PAT-INTEGRACAO-MOVEL',
            codigo='EQP-INTEGRACAO-MOVEL',
        )
        cls.operadora = OperadoraMovel.objects.create(
            empresa=cls.empresa,
            nome='Operadora integração',
        )
        cls.linha = LinhaMovel.objects.create(
            empresa=cls.empresa,
            base=cls.base,
            numero_normalizado='+5511987654321',
            operadora=cls.operadora,
            iccid='8955012345678901234',
            status=LinhaMovel.Status.EM_USO,
            criado_por=cls.admin,
        )
        cls.vinculo = VinculoLinhaEquipamento.objects.create(
            linha=cls.linha,
            equipamento=cls.equipamento,
            vinculado_por=cls.admin,
        )
        cls.token_protegido = 'TOKEN-CIFRADO-PIN-PUK-NAO-EXPORTAR'
        CredencialLinhaMovel.objects.create(
            linha=cls.linha,
            conteudo_criptografado=cls.token_protegido,
            chave_id='teste-integracao',
            campos_configurados=['pin', 'puk'],
            atualizado_por=cls.admin,
        )
        Historico.objects.create(
            equipamento=cls.equipamento,
            tipo_acao='CRIACAO',
            usuario=cls.admin,
        )

    def setUp(self):
        self.admin = User.objects.get(pk=self.admin.pk)
        self.admin_externo = User.objects.get(pk=self.admin_externo.pk)
        self.gestor_base = User.objects.get(pk=self.gestor_base.pk)
        self.client.force_login(self.admin)

    def test_api_publica_somente_numero_e_operadora_da_linha_autorizada(self):
        response = self.client.get(reverse(
            'estoque:equipamentos_por_regional',
            args=[self.produto.pk, self.base.pk],
        ))

        self.assertEqual(response.status_code, 200)
        equipamento = response.json()['equipamentos'][0]
        self.assertEqual(equipamento['conectividade_movel'], {
            'numero': self.linha.numero_formatado,
            'operadora': self.operadora.nome,
        })
        conteudo = response.content.decode()
        self.assertNotIn(self.linha.iccid, conteudo)
        self.assertNotIn(self.token_protegido, conteudo)
        self.assertNotIn('conteudo_criptografado', conteudo)

    def test_api_rejeita_base_de_outro_tenant_antes_de_buscar_equipamentos(self):
        response = self.client.get(reverse(
            'estoque:equipamentos_por_regional',
            args=[self.produto.pk, self.base_externa.pk],
        ))

        self.assertEqual(response.status_code, 404)

    def test_tory_exibe_linha_publica_apenas_em_consulta_direta_do_equipamento(self):
        resposta = AssistenteOperacionalService.responder(
            self.admin,
            f'Detalhe o equipamento de patrimônio {self.equipamento.patrimonio}',
        )['resposta']

        self.assertIn(self.linha.numero_formatado, resposta)
        self.assertIn(self.operadora.nome, resposta)
        self.assertNotIn(self.linha.iccid, resposta)
        self.assertNotIn(self.token_protegido, resposta)

        resumo = AssistenteOperacionalService.responder(
            self.admin,
            'Mostre todos os equipamentos',
        )['resposta']
        self.assertNotIn(self.linha.numero_formatado, resumo)

    def test_empresa_inativa_falha_fechada_na_projecao_e_na_tory(self):
        self.empresa.ativa = False
        self.empresa.save(update_fields=('ativa',))

        projecao = LinhasMoveisService.projecao_publica_por_equipamento(
            usuario=self.admin,
            equipamentos=[self.equipamento],
        )
        self.assertEqual(projecao, {})

        with self.assertRaises(PermissionDenied):
            AssistenteOperacionalService.responder(
                self.admin,
                f'Detalhe o equipamento de patrimônio {self.equipamento.patrimonio}',
            )

    def test_exportacao_excel_nao_inclui_identificadores_nem_credenciais_da_linha(self):
        response = self.client.get(reverse('estoque:exportar_historico_excel'))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        valores = '\n'.join(
            str(celula)
            for linha in workbook.active.iter_rows(values_only=True)
            for celula in linha
            if celula is not None
        )
        self.assertNotIn(self.linha.numero_normalizado, valores)
        self.assertNotIn(self.linha.iccid, valores)
        self.assertNotIn(self.token_protegido, valores)

    def test_notificacao_de_desvinculo_e_restrita_ao_tenant_e_sem_segredos(self):
        LinhasMoveisService.desvincular(
            usuario=self.admin,
            linha_id=self.linha.pk,
            motivo='Troca operacional planejada',
        )

        comunicados = list(Comunicado.objects.order_by('pk'))
        self.assertEqual(len(comunicados), 1)
        self.assertEqual(comunicados[0].titulo.lower(), 'linha móvel desvinculada')
        comunicado = comunicados[0]
        destinatarios = set(comunicado.usuarios.values_list('username', flat=True))
        self.assertIn(self.admin.username, destinatarios)
        self.assertNotIn(self.gestor_base.username, destinatarios)
        self.assertNotIn(self.admin_externo.username, destinatarios)
        conteudo = f'{comunicado.mensagem} {comunicado.dados}'
        self.assertNotIn(self.linha.iccid, conteudo)
        self.assertNotIn(self.token_protegido, conteudo)
        self.assertNotIn('pin', conteudo.lower())
        self.assertNotIn('puk', conteudo.lower())
