from datetime import timedelta
from io import BytesIO

import openpyxl
from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from estoque.models import (
    Base,
    Comunicado,
    Empresa,
    Equipamento,
    Modulo,
    ModuloEmpresa,
    Perfil,
    Produto,
    Sick,
    TermoEmpresa,
)
from estoque.policies.compras import GruposCorporativos
from estoque.services.assistente_operacional_service import (
    AssistenteOperacionalService,
)
from estoque.services.comunicado_service import ComunicadoService
from estoque.tenant_scope import TenantScope
from insumos.views.api import obter_base_importada


class MultitenantAuditStageITests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.company_a = Empresa.objects.create(nome='Tenant Auditoria I A')
        cls.company_b = Empresa.objects.create(nome='Tenant Auditoria I B')
        cls.inactive_company = Empresa.objects.create(
            nome='Tenant Auditoria I Inativa',
            ativa=False,
        )
        cls.base_a = Base.objects.create(nome='Unidade Auditoria A', empresa=cls.company_a)
        cls.base_b = Base.objects.create(nome='Unidade Auditoria B', empresa=cls.company_b)

        cls.admin_a = User.objects.create_user('audit.admin.a')
        cls.admin_a.perfil.role = Perfil.Role.ADMIN
        cls.admin_a.perfil.empresa = cls.company_a
        cls.admin_a.perfil.save()

        cls.admin_b = User.objects.create_user('audit.admin.b')
        cls.admin_b.perfil.role = Perfil.Role.ADMIN
        cls.admin_b.perfil.empresa = cls.company_b
        cls.admin_b.perfil.save()

        cls.superuser = User.objects.create_superuser(
            'audit.superuser',
            email='audit-superuser@example.test',
            password=None,
        )

        cls.operator_a = User.objects.create_user('audit.operator.a')
        cls.operator_a.perfil.role = Perfil.Role.OPERADOR
        cls.operator_a.perfil.empresa = cls.company_a
        cls.operator_a.perfil.save()
        cls.operator_a.perfil.regionais.add(cls.base_a)

        cls.operator_b = User.objects.create_user('audit.operator.b')
        cls.operator_b.perfil.role = Perfil.Role.OPERADOR
        cls.operator_b.perfil.empresa = cls.company_b
        cls.operator_b.perfil.save()
        cls.operator_b.perfil.regionais.add(cls.base_b)

        cls.product_a = Produto.objects.create(
            codigo='AUDIT-I-A',
            descricao='Máquina A',
            fabricante='Tenant A',
            modelo='A',
            categoria='Categoria exclusiva A',
            empresa_catalogo_origem=cls.company_a,
        )
        cls.product_b = Produto.objects.create(
            codigo='AUDIT-I-B',
            descricao='Máquina B',
            fabricante='Tenant B',
            modelo='B',
            categoria='Categoria exclusiva B',
            empresa_catalogo_origem=cls.company_b,
        )
        cls.equipment_a = Equipamento.objects.create(
            produto=cls.product_a,
            regional=cls.base_a,
            numero_serie='AUDIT-SERIE-A',
            patrimonio='AUDIT-PAT-A',
            codigo='AUDIT-EQP-A',
            status='MANUTENCAO',
        )
        cls.equipment_b = Equipamento.objects.create(
            produto=cls.product_b,
            regional=cls.base_b,
            numero_serie='AUDIT-SERIE-B',
            patrimonio='AUDIT-PAT-B',
            codigo='AUDIT-EQP-B',
            status='MANUTENCAO',
        )
        cls.sick_a = Sick.objects.create(
            equipamento=cls.equipment_a,
            categoria='HARDWARE',
            motivo='Auditoria A',
            previsao_retorno=timezone.localdate() + timedelta(days=1),
            status_final='MANUTENCAO',
            ativo=True,
        )
        cls.sick_b = Sick.objects.create(
            equipamento=cls.equipment_b,
            categoria='HARDWARE',
            motivo='Auditoria B',
            previsao_retorno=timezone.localdate() + timedelta(days=1),
            status_final='MANUTENCAO',
            ativo=True,
        )

        maintenance_group, _ = Group.objects.get_or_create(
            name=GruposCorporativos.SICK_MANUTENCAO,
        )
        cls.operator_a.groups.add(maintenance_group)
        cls.operator_b.groups.add(maintenance_group)

        sick_module, _ = Modulo.objects.get_or_create(
            codigo=Modulo.Codigo.SICK,
            defaults={'nome': 'SICK'},
        )
        ModuloEmpresa.objects.update_or_create(
            empresa=cls.company_a,
            modulo=sick_module,
            defaults={'habilitado': True, 'nome_exibicao': 'Manutenção'},
        )
        TermoEmpresa.objects.create(
            empresa=cls.company_a,
            chave=TermoEmpresa.Chave.EQUIPAMENTO,
            valor_singular='Máquina',
            valor_plural='Máquinas',
        )
        TermoEmpresa.objects.create(
            empresa=cls.company_a,
            chave=TermoEmpresa.Chave.BASE,
            valor_singular='Unidade',
            valor_plural='Unidades',
        )
        TermoEmpresa.objects.create(
            empresa=cls.company_a,
            chave=TermoEmpresa.Chave.REGIONAL,
            valor_singular='Território',
            valor_plural='Territórios',
        )

    def test_company_notification_never_reaches_foreign_admin(self):
        communication = ComunicadoService.criar_acao(
            titulo='Alerta privado A',
            mensagem='Conteúdo exclusivo do tenant A',
            usuario=self.operator_a,
            empresa=self.company_a,
            incluir_autor=False,
        )

        recipients = set(
            communication.usuarios.values_list('username', flat=True)
        )
        self.assertSetEqual(recipients, {'audit.admin.a', 'audit.superuser'})
        self.assertNotIn('audit.admin.b', recipients)

    def test_base_notification_scopes_admins_by_company(self):
        recipients = set(
            ComunicadoService.usuarios_por_bases([self.base_a]).values_list(
                'username', flat=True,
            )
        )

        self.assertSetEqual(
            recipients,
            {'audit.operator.a', 'audit.admin.a', 'audit.superuser'},
        )

    def test_scheduled_maintenance_notifications_are_tenant_scoped(self):
        communications = ComunicadoService.notificar_manutencoes_previstas()
        by_company = {item.empresa_id: item for item in communications}

        recipients_a = set(
            by_company[self.company_a.pk].usuarios.values_list(
                'username', flat=True,
            )
        )
        recipients_b = set(
            by_company[self.company_b.pk].usuarios.values_list(
                'username', flat=True,
            )
        )
        self.assertSetEqual(
            recipients_a,
            {'audit.operator.a', 'audit.admin.a', 'audit.superuser'},
        )
        self.assertSetEqual(
            recipients_b,
            {'audit.operator.b', 'audit.admin.b', 'audit.superuser'},
        )

    def test_sick_filters_do_not_disclose_foreign_tenant(self):
        self.client.force_login(self.admin_a)

        response = self.client.get(reverse('estoque:sick'))

        self.assertEqual(response.status_code, 200)
        self.assertQuerySetEqual(
            response.context['regionais'],
            [self.base_a],
            transform=lambda item: item,
        )
        self.assertEqual(list(response.context['categorias']), ['Categoria exclusiva A'])
        self.assertNotContains(response, 'Categoria exclusiva B')
        self.assertNotContains(response, 'Unidade Auditoria B')

    def test_tory_output_uses_tenant_terminology_without_changing_codes(self):
        scope = TenantScope.fresh_for_user(self.admin_a)
        AssistenteOperacionalService._authorize_tenant_scope(
            self.admin_a,
            scope,
        )

        response = AssistenteOperacionalService._aplicar_terminologia_tenant(
            self.admin_a,
            {
                'categoria': 'estoque',
                'resposta': 'Equipamentos em SICK por base.',
                'acoes': [
                    {'label': 'Ver equipamentos', 'pergunta': 'Equipamentos na base A'},
                ],
            },
        )

        self.assertEqual(response['categoria'], 'estoque')
        self.assertEqual(response['resposta'], 'Máquinas em Manutenção por unidade.')
        self.assertEqual(response['acoes'][0]['label'], 'Ver máquinas')
        self.assertEqual(response['acoes'][0]['pergunta'], 'Máquinas na unidade A')

    def test_excel_import_selector_only_lists_active_companies(self):
        self.client.force_login(self.superuser)

        response = self.client.get(reverse('insumos:importar_excel'))

        self.assertEqual(response.status_code, 200)
        self.assertQuerySetEqual(
            response.context['empresas_importacao'],
            [self.company_a, self.company_b],
            transform=lambda item: item,
            ordered=False,
        )
        self.assertContains(response, 'Empresa principal')
        self.assertNotContains(response, self.inactive_company.nome)

    def test_excel_import_fails_closed_without_explicit_company(self):
        self.client.force_login(self.superuser)
        workbook = openpyxl.Workbook()
        buffer = BytesIO()
        workbook.save(buffer)
        upload = SimpleUploadedFile(
            'auditoria.xlsx',
            buffer.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            ),
        )

        response = self.client.post(
            reverse('insumos:importar_excel'),
            {'arquivo': upload},
            follow=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Selecione a empresa principal da importação.')

    def test_excel_import_never_infers_alternate_company_by_name(self):
        coincidental_company = Empresa.objects.create(nome='OXXO')
        explicit_company = Empresa.objects.create(nome='Alternativa explícita')
        coincidental_base = Base.objects.create(
            nome='OXXO REGIONAL X',
            empresa=coincidental_company,
        )
        explicit_base = Base.objects.create(
            nome='OXXO REGIONAL X',
            empresa=explicit_company,
        )

        without_alternate = obter_base_importada(
            'REGIONAL X',
            {},
            self.company_a,
        )
        with_alternate = obter_base_importada(
            'REGIONAL X',
            {},
            self.company_a,
            explicit_company,
        )

        self.assertIsNone(without_alternate)
        self.assertEqual(with_alternate, explicit_base)
        self.assertNotEqual(with_alternate, coincidental_base)

    def test_user_facing_exports_use_tenant_terminology(self):
        self.client.force_login(self.admin_a)

        inventory_response = self.client.get(
            reverse('estoque:exportar_historico_excel'),
        )
        inventory_workbook = openpyxl.load_workbook(
            BytesIO(inventory_response.content),
            read_only=True,
        )
        inventory_headers = [
            cell.value
            for cell in next(inventory_workbook.active.iter_rows(max_row=1))
        ]

        tickets_response = self.client.get(reverse('chamados:exportar'))
        tickets_workbook = openpyxl.load_workbook(
            BytesIO(b''.join(tickets_response.streaming_content)),
            read_only=True,
        )
        ticket_headers = [
            cell.value
            for cell in next(tickets_workbook.active.iter_rows(max_row=1))
        ]

        self.assertEqual(inventory_response.status_code, 200)
        self.assertIn('Território', inventory_headers)
        self.assertIn('Máquina', inventory_headers)
        self.assertEqual(tickets_response.status_code, 200)
        self.assertIn('TERRITÓRIO', ticket_headers)

