from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from estoque.models import (
    Base,
    CapacidadeRelacionamentoEmpresa,
    Empresa,
    Equipamento,
    Historico,
    Perfil,
    Produto,
    RelacionamentoEmpresa,
)
from estoque.security import (
    secure_base_queryset,
    secure_history_queryset,
    secure_queryset,
)


class InventoryTenantIsolationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.brasil = Empresa.objects.create(nome='Inventory Brasil isolamento')
        cls.latam = Empresa.objects.create(nome='Inventory LATAM isolamento')
        cls.externa = Empresa.objects.create(nome='Empresa externa isolamento')
        cls.base_brasil = Base.objects.create(nome='Base Brasil isolamento', empresa=cls.brasil)
        cls.base_latam = Base.objects.create(nome='Base LATAM isolamento', empresa=cls.latam)
        cls.base_externa = Base.objects.create(nome='Base externa isolamento', empresa=cls.externa)
        cls.produto = Produto.objects.create(
            codigo='PROD-ISOLAMENTO-ETAPA13',
            descricao='Produto isolamento Etapa 13',
            fabricante='Inventory',
            modelo='Tenant',
            categoria='Coletores',
        )
        cls.equip_brasil = cls._equipment(cls.base_brasil, 'BR')
        cls.equip_latam = cls._equipment(cls.base_latam, 'LATAM')
        cls.equip_externo = cls._equipment(cls.base_externa, 'EXT')
        cls.admin = User.objects.create_user(
            'etapa13.admin', password='SenhaTeste123!'
        )
        cls.admin.perfil.role = Perfil.Role.ADMIN
        cls.admin.perfil.empresa = cls.brasil
        cls.admin.perfil.save()

        cls.admin_restrito = User.objects.create_user(
            'etapa13.admin.restrito', password='SenhaTeste123!'
        )
        cls.admin_restrito.perfil.role = Perfil.Role.ADMIN
        cls.admin_restrito.perfil.empresa = cls.brasil
        cls.admin_restrito.perfil.save()
        cls.hist_brasil = Historico.objects.create(
            equipamento=cls.equip_brasil,
            tipo_acao='CRIACAO',
            usuario=cls.admin,
        )
        cls.hist_externo = Historico.objects.create(
            equipamento=cls.equip_externo,
            tipo_acao='CRIACAO',
            usuario=cls.admin,
        )

        cls.relationship = RelacionamentoEmpresa.objects.create(
            empresa_origem=cls.brasil,
            empresa_destino=cls.latam,
        )
        CapacidadeRelacionamentoEmpresa.objects.create(
            relacionamento=cls.relationship,
            recurso=CapacidadeRelacionamentoEmpresa.Recurso.EQUIPAMENTOS,
            acao=CapacidadeRelacionamentoEmpresa.Acao.VISUALIZAR,
        )
        cls.admin.perfil.empresas_acesso_adicional.add(cls.latam)

    @classmethod
    def _equipment(cls, base, suffix):
        return Equipamento.objects.create(
            produto=cls.produto,
            numero_serie=f'SERIE-ETAPA13-{suffix}',
            patrimonio=f'PATRIMONIO-ETAPA13-{suffix}',
            regional=base,
            status='ATIVO',
            codigo=f'EQP-ETAPA13-{suffix}',
        )

    def test_querysets_apply_selected_companies_capability_and_bases(self):
        equipment_ids = set(
            secure_queryset(Equipamento.objects.all(), self.admin)
            .values_list('pk', flat=True)
        )
        base_ids = set(
            secure_base_queryset(Base.objects.all(), self.admin)
            .values_list('pk', flat=True)
        )
        history_ids = set(
            secure_history_queryset(Historico.objects.all(), self.admin)
            .values_list('pk', flat=True)
        )

        self.assertEqual(equipment_ids, {self.equip_brasil.pk, self.equip_latam.pk})
        self.assertEqual(base_ids, {self.base_brasil.pk, self.base_latam.pk})
        self.assertIn(self.hist_brasil.pk, history_ids)
        self.assertNotIn(self.hist_externo.pk, history_ids)

        restricted_ids = set(
            secure_queryset(Equipamento.objects.all(), self.admin_restrito)
            .values_list('pk', flat=True)
        )
        self.assertEqual(restricted_ids, {self.equip_brasil.pk})

    def test_stock_ajax_lists_selected_related_but_never_external_company(self):
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse('estoque:detalhes_produto', args=[self.produto.pk])
        )

        self.assertEqual(response.status_code, 200)
        returned_ids = {item['id'] for item in response.json()['equipamentos']}
        self.assertEqual(returned_ids, {self.equip_brasil.pk, self.equip_latam.pk})
        self.assertNotIn(self.equip_externo.pk, returned_ids)

    def test_quick_search_finds_only_equipment_in_visible_tenants(self):
        self.client.force_login(self.admin)
        url = reverse('estoque:busca_rapida_equipamento_api')

        for field, value, expected_type in (
            ('numero_serie', self.equip_brasil.numero_serie, 'numero_serie'),
            ('patrimonio', self.equip_latam.patrimonio, 'patrimonio'),
        ):
            with self.subTest(field=field):
                response = self.client.get(url, {'q': value})

                self.assertEqual(response.status_code, 200)
                payload = response.json()
                self.assertTrue(payload['encontrado'])
                self.assertEqual(payload['tipo'], expected_type)
                self.assertEqual(payload['equipamento'][field], value)

        external_response = self.client.get(
            url,
            {'q': self.equip_externo.numero_serie},
        )

        self.assertEqual(external_response.status_code, 200)
        self.assertEqual(external_response.json(), {'encontrado': False})

    def test_quick_search_rejects_short_terms_without_exposing_data(self):
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse('estoque:busca_rapida_equipamento_api'),
            {'q': 'A'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'encontrado': False})

    def test_selected_company_without_equipment_capability_is_not_visible(self):
        self.relationship.capacidades.all().delete()

        equipment_ids = set(
            secure_queryset(Equipamento.objects.all(), self.admin)
            .values_list('pk', flat=True)
        )

        self.assertEqual(equipment_ids, {self.equip_brasil.pk})

    def test_manual_ids_cannot_open_external_base_equipment_or_history(self):
        self.client.force_login(self.admin)
        urls = (
            reverse('estoque:detalhes_regional_api', args=[self.base_externa.pk]),
            reverse(
                'estoque:equipamentos_por_regional',
                args=[self.produto.pk, self.base_externa.pk],
            ),
            reverse('estoque:editar_equipamento', args=[self.equip_externo.pk]),
            reverse('estoque:historico_detalhes', args=[self.hist_externo.pk]),
            reverse('estoque:historico_modal', args=[self.equip_externo.pk]),
        )

        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 404)

    def test_manual_filter_ids_cannot_select_external_company_or_base(self):
        self.client.force_login(self.admin)

        self.assertEqual(
            self.client.get(
                reverse('estoque:index'), {'inventory': self.externa.pk}
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.get(
                reverse('estoque:detalhes_produto', args=[self.produto.pk]),
                {'regional': self.base_externa.pk},
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.get(
                reverse('estoque:api_equipamentos'),
                {'regional': self.base_externa.pk, 'categoria': 'Coletores'},
            ).status_code,
            404,
        )

    def test_create_rejects_forged_external_base(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('estoque:cadastrar_equipamento'),
            {
                'categoria': 'Coletores',
                'produto': str(self.produto.pk),
                'numero_serie': 'SERIE-ETAPA13-FORJADA',
                'patrimonio': 'PATRIMONIO-ETAPA13-FORJADO',
                'regional': str(self.base_externa.pk),
                'finalidade': Equipamento.Finalidade.OPERACIONAL,
                'responsavel': '',
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(
            Equipamento.objects.filter(numero_serie='SERIE-ETAPA13-FORJADA').exists()
        )

    def test_related_edit_requires_exact_capability(self):
        self.client.force_login(self.admin)
        url = reverse('estoque:editar_equipamento', args=[self.equip_latam.pk])

        self.assertEqual(self.client.get(url).status_code, 404)

        CapacidadeRelacionamentoEmpresa.objects.create(
            relacionamento=self.relationship,
            recurso=CapacidadeRelacionamentoEmpresa.Recurso.EQUIPAMENTOS,
            acao=CapacidadeRelacionamentoEmpresa.Acao.EDITAR,
        )
        self.assertEqual(self.client.get(url).status_code, 200)

    def test_excel_export_contains_only_visible_tenants(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse('estoque:exportar_historico_excel'))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        values = {
            str(cell)
            for row in workbook.active.iter_rows(values_only=True)
            for cell in row
            if cell is not None
        }
        self.assertIn(self.equip_brasil.numero_serie, values)
        self.assertIn(self.equip_latam.numero_serie, values)
        self.assertNotIn(self.equip_externo.numero_serie, values)
