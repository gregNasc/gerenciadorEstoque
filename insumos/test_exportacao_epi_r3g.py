from datetime import timedelta
from io import BytesIO

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from estoque.models import Base, Empresa, Perfil
from insumos.models import (
    CategoriaInsumo,
    CertificadoEpi,
    Insumo,
    LoteEpi,
    MovimentacaoInsumo,
    ProdutoEpi,
    VarianteEpi,
)


class ExportacaoEpiR3GTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.empresa = Empresa.objects.create(nome='Empresa exportação EPI')
        cls.base = Base.objects.create(nome='Base exportação EPI', empresa=cls.empresa)
        cls.outra_empresa = Empresa.objects.create(nome='Outro tenant EPI')
        cls.outra_base = Base.objects.create(nome='Base proibida EPI', empresa=cls.outra_empresa)
        cls.admin = User.objects.create_user('admin.exportacao.epi', password='senha')
        cls.admin.perfil.empresa = cls.empresa
        cls.admin.perfil.role = Perfil.Role.ADMIN
        cls.admin.perfil.save(update_fields=['empresa', 'role'])
        categoria = CategoriaInsumo.objects.create(nome='EPI')
        cls.insumo = Insumo.objects.create(
            descricao='Botas exportação',
            categoria=categoria,
            unidade_medida='PAR',
            natureza=Insumo.Natureza.EPI,
        )
        produto = ProdutoEpi.objects.create(insumo=cls.insumo)
        variante = VarianteEpi.objects.create(
            produto_epi=produto,
            nome='Botas',
            tamanho='37',
        )
        hoje = timezone.localdate()
        vencido = CertificadoEpi.objects.create(
            produto_epi=produto,
            numero_ca='11111',
            validade_ca=hoje - timedelta(days=1),
            situacao_oficial='VENCIDO',
            snapshot_oficial={'natureza': 'Nacional'},
        )
        proximo = CertificadoEpi.objects.create(
            produto_epi=produto,
            numero_ca='22222',
            validade_ca=hoje + timedelta(days=90),
            situacao_oficial='VÁLIDO',
        )
        vigente = CertificadoEpi.objects.create(
            produto_epi=produto,
            numero_ca='33333',
            validade_ca=hoje + timedelta(days=400),
            situacao_oficial='VÁLIDO',
        )
        lote_vencido = LoteEpi.objects.create(
            base=cls.base, produto_epi=produto, variante=variante,
            certificado=vencido, codigo_lote='VENCIDO-1', ativo=True,
        )
        lote_proximo = LoteEpi.objects.create(
            base=cls.base, produto_epi=produto, variante=variante,
            certificado=proximo, codigo_lote='PROXIMO-1', ativo=False,
        )
        lote_vigente = LoteEpi.objects.create(
            base=cls.base, produto_epi=produto, variante=variante,
            certificado=vigente, codigo_lote='VIGENTE-1', ativo=True,
        )
        for lote, quantidade in (
            (lote_vencido, 2), (lote_proximo, 3), (lote_vigente, 4),
        ):
            MovimentacaoInsumo.objects.create(
                base=cls.base,
                insumo=cls.insumo,
                tipo='ENTRADA',
                quantidade=quantidade,
                lote_epi=lote,
                usuario=cls.admin,
            )

    def test_exporta_filtros_de_validade_e_atividade_sem_precos(self):
        self.client.force_login(self.admin)
        resposta = self.client.get(reverse('insumos:exportar_dados_epi'), {
            'base': self.base.pk,
            'insumo': self.insumo.pk,
            'validade': ['vencido', 'proximo'],
            'registro': ['ativo', 'inativo'],
        })

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(
            resposta['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        planilha = openpyxl.load_workbook(BytesIO(resposta.content)).active
        valores = [celula for linha in planilha.iter_rows(values_only=True) for celula in linha]
        self.assertIn('VENCIDO-1', valores)
        self.assertIn('PROXIMO-1', valores)
        self.assertNotIn('VIGENTE-1', valores)
        cabecalhos = [celula.value for celula in planilha[1]]
        self.assertNotIn('PREÇO', cabecalhos)
        self.assertNotIn('CUSTO', cabecalhos)
        self.assertNotIn('VALOR', cabecalhos)
        self.assertIn('QUANTIDADE', cabecalhos)
        indice_lote = cabecalhos.index('LOTE')
        indice_quantidade = cabecalhos.index('QUANTIDADE')
        quantidades = {
            linha[indice_lote].value: linha[indice_quantidade].value
            for linha in planilha.iter_rows(min_row=2)
        }
        self.assertEqual(quantidades, {'VENCIDO-1': 2, 'PROXIMO-1': 3})

    def test_rejeita_exportacao_de_base_de_outro_tenant(self):
        self.client.force_login(self.admin)
        resposta = self.client.get(reverse('insumos:exportar_dados_epi'), {
            'base': self.outra_base.pk,
            'insumo': self.insumo.pk,
        })
        self.assertEqual(resposta.status_code, 403)

    def test_sem_filtros_explicitos_exporta_todas_as_faixas(self):
        self.client.force_login(self.admin)

        resposta = self.client.get(reverse('insumos:exportar_dados_epi'), {
            'base': self.base.pk,
            'insumo': self.insumo.pk,
        })

        self.assertEqual(resposta.status_code, 200)
        planilha = openpyxl.load_workbook(BytesIO(resposta.content)).active
        lotes = {
            linha[15].value
            for linha in planilha.iter_rows(min_row=2)
        }
        self.assertEqual(lotes, {'VENCIDO-1', 'PROXIMO-1', 'VIGENTE-1'})

